"""Excel and PDF exports from an AnalyticsReport (same numbers as the dashboard)."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.utils.translation import gettext as _

from .services import AnalyticsReport


def _money(value: Decimal) -> str:
    return f"{value.quantize(Decimal('0.01'))}"


def _pct(value: Decimal | None) -> str:
    if value is None:
        return "—"
    return f"{value.quantize(Decimal('0.01'))} %"


def _delta_line(label: str, delta) -> tuple[str, str, str, str, str]:
    return (
        label,
        _money(delta.current),
        _money(delta.previous),
        _money(delta.difference),
        _pct(delta.pct),
    )


def build_xlsx(report: AnalyticsReport) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    period = report.filters.period
    header_font = Font(bold=True)

    kpi = wb.active
    kpi.title = _("Indicateurs")
    kpi.append([_("Restaurant"), report.restaurant_name])
    kpi.append([_("Période"), period.label()])
    kpi.append([_("Du"), period.start_date.isoformat()])
    kpi.append([_("Au"), period.end_date.isoformat()])
    kpi.append([_("Devise"), report.currency])
    kpi.append([_("Généré le"), report.generated_at.strftime("%Y-%m-%d %H:%M")])
    kpi.append([])
    kpi.append([_("Indicateur"), _("Actuel"), _("Précédent"), _("Écart"), _("Évolution")])
    for cell in kpi[kpi.max_row]:
        cell.font = header_font
    for row in (
        _delta_line(_("CA payé"), report.revenue),
        _delta_line(_("Commandes"), report.orders_total),
        _delta_line(_("Commandes payées"), report.orders_paid),
        _delta_line(_("Commandes annulées"), report.orders_cancelled),
        _delta_line(_("Panier moyen"), report.avg_ticket),
        _delta_line(_("Taux d'annulation"), report.cancel_rate),
    ):
        kpi.append(row)
    kpi.append([_("Commandes en cours"), report.orders_open])
    kpi.append([_("CA brut"), _money(report.gross_revenue)])
    kpi.append([_("CA net"), _money(report.net_revenue)])
    kpi.append([_("Remises (nombre)"), report.discount_count])
    kpi.append([_("Remises approuvées"), report.discount_approved_count, _money(report.discount_approved_amount)])
    kpi.append([_("Remises refusées"), report.discount_rejected_count, _money(report.discount_rejected_amount)])
    kpi.append([_("Remises en attente"), report.discount_pending_count])

    dishes = wb.create_sheet(_("Plats"))
    dishes.append([_("Plat"), _("Quantité"), _("CA"), _("% des ventes")])
    for cell in dishes[1]:
        cell.font = header_font
    for row in report.top_dishes:
        dishes.append([row.name, row.quantity, _money(row.revenue), _pct(row.share_pct)])

    least = wb.create_sheet(_("Plats les moins vendus"))
    least.append([_("Plat"), _("Quantité"), _("CA")])
    for cell in least[1]:
        cell.font = header_font
    for row in report.bottom_dishes:
        least.append([row.name, row.quantity, _money(row.revenue)])

    cats = wb.create_sheet(_("Catégories"))
    cats.append([_("Catégorie"), _("Ventes"), _("CA"), _("% du CA")])
    for cell in cats[1]:
        cell.font = header_font
    for row in report.categories:
        cats.append([row.name, row.quantity, _money(row.revenue), _pct(row.share_pct)])

    tables = wb.create_sheet(_("Tables"))
    tables.append([_("Table"), _("Commandes"), _("CA")])
    for cell in tables[1]:
        cell.font = header_font
    for row in report.tables:
        tables.append([row.number, row.orders, _money(row.revenue)])

    resa = wb.create_sheet(_("Réservations"))
    r = report.reservations
    resa.append([_("Total"), r.total])
    resa.append([_("En attente"), r.pending])
    resa.append([_("Confirmées"), r.confirmed])
    resa.append([_("Refusées"), r.refused])
    resa.append([_("Annulées"), r.cancelled])
    resa.append([_("Personnes"), r.guests])
    resa.append([_("Taux de confirmation"), _pct(r.confirm_rate)])
    resa.append([_("Taux d'annulation"), _pct(r.cancel_rate)])

    cash = wb.create_sheet(_("Caisse"))
    c = report.cash
    cash.append([_("Entrées"), _money(c.total_in)])
    cash.append([_("Entrées commandes"), _money(c.order_in)])
    cash.append([_("Entrées manuelles"), _money(c.manual_in)])
    cash.append([_("Sorties"), _money(c.total_out)])
    cash.append([_("Solde"), _money(c.balance)])
    cash.append([_("Clôtures"), c.closures])
    cash.append([_("Somme des écarts de clôture"), _money(c.closure_gap_sum)])
    cash.append([_("Écart CA payé vs encaissements commandes"), _money(c.vs_revenue_gap)])

    rem = wb.create_sheet(_("Remises"))
    rem.append([_("Demandées par"), _("Nombre"), _("Montant")])
    for cell in rem[1]:
        cell.font = header_font
    for row in report.discounts_by_requester:
        rem.append([row.name, row.count, _money(row.amount)])

    insights = wb.create_sheet(_("Observations"))
    if report.insights:
        for line in report.insights:
            insights.append([line])
    else:
        insights.append([_("Aucune observation calculable pour cette période.")])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(report: AnalyticsReport) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=_("Analyse"),
    )
    styles = getSampleStyleSheet()
    story = []
    period = report.filters.period
    cur = report.currency

    story.append(Paragraph(f"{report.restaurant_name} — {_('Analyse')}", styles["Title"]))
    story.append(
        Paragraph(
            f"{period.label()} · {period.start_date:%d/%m/%Y} – {period.end_date:%d/%m/%Y} · {cur}",
            styles["Normal"],
        )
    )
    story.append(
        Paragraph(
            f"{_('Généré le')} {report.generated_at.strftime('%d/%m/%Y %H:%M')}",
            styles["Normal"],
        )
    )
    story.append(Spacer(1, 8 * mm))

    kpi_data = [
        [_("Indicateur"), _("Actuel"), _("Précédent"), _("Évolution")],
        [_("CA payé"), f"{_money(report.revenue.current)} {cur}", f"{_money(report.revenue.previous)} {cur}", _pct(report.revenue.pct)],
        [_("Commandes"), str(int(report.orders_total.current)), str(int(report.orders_total.previous)), _pct(report.orders_total.pct)],
        [_("Payées"), str(int(report.orders_paid.current)), str(int(report.orders_paid.previous)), _pct(report.orders_paid.pct)],
        [_("Annulées"), str(int(report.orders_cancelled.current)), str(int(report.orders_cancelled.previous)), _pct(report.orders_cancelled.pct)],
        [_("Panier moyen"), f"{_money(report.avg_ticket.current)} {cur}", f"{_money(report.avg_ticket.previous)} {cur}", _pct(report.avg_ticket.pct)],
        [_("Taux d'annulation"), _pct(report.cancel_rate.current), _pct(report.cancel_rate.previous), _pct(report.cancel_rate.pct)],
    ]
    story.append(_styled_table(kpi_data))
    story.append(Spacer(1, 6 * mm))

    if report.insights:
        story.append(Paragraph(_("Analyse intelligente"), styles["Heading2"]))
        for line in report.insights:
            story.append(Paragraph(f"• {line}", styles["Normal"]))
        story.append(Spacer(1, 4 * mm))

    if report.top_dishes:
        story.append(Paragraph(_("Plats les plus vendus"), styles["Heading2"]))
        data = [[_("Plat"), _("Quantité"), _("CA"), _("%")]]
        for row in report.top_dishes:
            data.append([row.name, str(row.quantity), f"{_money(row.revenue)} {cur}", _pct(row.share_pct)])
        story.append(_styled_table(data))
        story.append(Spacer(1, 4 * mm))

    if report.categories:
        story.append(Paragraph(_("Catégories"), styles["Heading2"]))
        data = [[_("Catégorie"), _("Ventes"), _("CA"), _("%")]]
        for row in report.categories:
            data.append([row.name, str(row.quantity), f"{_money(row.revenue)} {cur}", _pct(row.share_pct)])
        story.append(_styled_table(data))
        story.append(Spacer(1, 4 * mm))

    if report.tables:
        story.append(Paragraph(_("Tables"), styles["Heading2"]))
        data = [[_("Table"), _("Commandes"), _("CA")]]
        for row in report.tables:
            data.append([str(row.number), str(row.orders), f"{_money(row.revenue)} {cur}"])
        story.append(_styled_table(data))
        story.append(Spacer(1, 4 * mm))

    r = report.reservations
    story.append(Paragraph(_("Réservations"), styles["Heading2"]))
    story.append(
        _styled_table(
            [
                [_("Total"), _("Confirmées"), _("Refusées"), _("Annulées"), _("Personnes")],
                [str(r.total), str(r.confirmed), str(r.refused), str(r.cancelled), str(r.guests)],
            ]
        )
    )
    story.append(Spacer(1, 4 * mm))

    c = report.cash
    story.append(Paragraph(_("Caisse"), styles["Heading2"]))
    story.append(
        _styled_table(
            [
                [_("Entrées"), _("Dont commandes"), _("Manuelles"), _("Sorties"), _("Écart vs CA")],
                [
                    f"{_money(c.total_in)} {cur}",
                    f"{_money(c.order_in)} {cur}",
                    f"{_money(c.manual_in)} {cur}",
                    f"{_money(c.total_out)} {cur}",
                    f"{_money(c.vs_revenue_gap)} {cur}",
                ],
            ]
        )
    )

    if report.revenue_series:
        story.append(Spacer(1, 6 * mm))
        story.append(Paragraph(_("Évolution du CA"), styles["Heading2"]))
        story.append(_pdf_bar_chart(report.revenue_series, report.currency))

    if any(n for _h, n in report.hourly):
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph(_("Commandes par heure"), styles["Heading2"]))
        story.append(_pdf_bar_chart([(f"{h:02d}h", Decimal(n)) for h, n in report.hourly if n], ""))

    doc.build(story)
    return buf.getvalue()


def _styled_table(data):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A0A0A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#E4E4E7")),
                ("BACKGROUND", (0, 1), (-1, -1), colors.white),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def _pdf_bar_chart(series: list[tuple[str, Decimal]], currency: str):
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors

    values = [float(v) for _label, v in series]
    labels = [label for label, _v in series]
    width = min(480, max(260, 28 * len(values)))
    drawing = Drawing(width, 160)
    chart = VerticalBarChart()
    chart.x = 30
    chart.y = 25
    chart.height = 120
    chart.width = width - 50
    chart.data = [values]
    chart.categoryAxis.categoryNames = labels
    chart.categoryAxis.labels.angle = 45
    chart.categoryAxis.labels.fontSize = 6
    chart.barWidth = 8
    chart.groupSpacing = 8
    chart.bars[0].fillColor = colors.HexColor("#FF5A00")
    chart.valueAxis.labels.fontSize = 7
    drawing.add(chart)
    return drawing


def xlsx_response(report: AnalyticsReport) -> HttpResponse:
    period = report.filters.period
    filename = f"analytics-{period.start_date.isoformat()}-{period.end_date.isoformat()}.xlsx"
    response = HttpResponse(
        build_xlsx(report),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def pdf_response(report: AnalyticsReport) -> HttpResponse:
    period = report.filters.period
    filename = f"analytics-{period.start_date.isoformat()}-{period.end_date.isoformat()}.pdf"
    response = HttpResponse(build_pdf(report), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
